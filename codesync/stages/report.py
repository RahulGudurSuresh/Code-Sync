"""Findings -> summary numbers, a Markdown report and an Excel workbook."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..schema import DocumentMeta, Finding, Requirement


def summarize(findings: list[Finding], reqs_a: list[Requirement],
              reqs_b: list[Requirement]) -> dict:
    def count(**want) -> int:
        return sum(all(getattr(f, k) == v for k, v in want.items())
                   for f in findings)
    conflicts = [f for f in findings if f.category == "conflict"]
    return {
        "requirements_a": sum(not r.excluded for r in reqs_a),
        "requirements_b": sum(not r.excluded for r in reqs_b),
        "set_aside_a": sum(r.excluded for r in reqs_a),
        "set_aside_b": sum(r.excluded for r in reqs_b),
        "conflicts": len(conflicts),
        "contradictions": count(type="contradiction"),
        "more_stringent": count(type="more_stringent"),
        "similar": count(category="similar"),
        "unique_a": count(type="only_in_a"),
        "unique_b": count(type="only_in_b"),
        "severity": {s: sum(f.severity == s for f in conflicts)
                     for s in ("high", "medium", "low")},
        "needs_review": sum(f.needs_review for f in findings),
    }


def _cite(r: Requirement | None, also: list[str] | None = None) -> str:
    if r is None:
        return ""
    src = "drawing" if r.from_figure else r.modality
    text = f"**{r.ref}** (p.{r.page}, {src}): {r.statement}"
    if also:
        text += " *Also stated at: " + "; ".join(also) + ".*"
    return text


def _cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def markdown_report(meta_a: DocumentMeta, meta_b: DocumentMeta,
                    governing: DocumentMeta, findings: list[Finding],
                    summary: dict, requirements: list[Requirement]) -> str:
    a, b = meta_a.label, meta_b.label
    lines = [
        f"# Code comparison: {a} vs {b}", "",
        f"- **{a}**: {meta_a.name} ({meta_a.filename}, {meta_a.pages} pages)",
        f"- **{b}**: {meta_b.name} ({meta_b.filename}, {meta_b.pages} pages)",
        f"- **Governing document**: {governing.label}", "",
        "## Summary", "",
        "| | Count |", "|---|---|",
        f"| Requirements in {a} | {summary['requirements_a']} |",
        f"| Requirements in {b} | {summary['requirements_b']} |",
        f"| Conflicts: direct contradictions | {summary['contradictions']} |",
        f"| Conflicts: one code stricter | {summary['more_stringent']} |",
        f"| Similar (equivalent) requirements | {summary['similar']} |",
        f"| Only in {a} | {summary['unique_a']} |",
        f"| Only in {b} | {summary['unique_b']} |",
        f"| Findings flagged for engineer review | {summary['needs_review']} |",
        f"| Drawing statements set aside for manual reading | "
        f"{summary['set_aside_a'] + summary['set_aside_b']} |",
        "",
    ]

    def conflict_table(title: str, kind: str) -> None:
        rows = [f for f in findings if f.type == kind]
        lines.extend([f"## {title} ({len(rows)})", ""])
        if not rows:
            lines.extend(["None found.", ""])
            return
        lines.extend([f"| ID | Topic | Severity | {a} | {b} | Action | Review |",
                      "|---|---|---|---|---|---|---|"])
        for f in rows:
            review = f.review_reason if f.needs_review else ""
            if f.drawing_check == "confirmed":
                review = (review + " " if review else "") + \
                    "Drawing value confirmed."
            lines.append("| " + " | ".join(_cell(c) for c in (
                f.id, f.topic, f.severity, _cite(f.a), _cite(f.b, f.also_b),
                f.action, review)) + " |")
        lines.append("")

    conflict_table("Direct contradictions", "contradiction")
    conflict_table("One code stricter than the other", "more_stringent")

    similar = [f for f in findings if f.category == "similar"]
    lines.extend([f"## Similar requirements ({len(similar)})", ""])
    if similar:
        lines.extend([f"| ID | {a} | {b} |", "|---|---|---|"])
        lines.extend(f"| {f.id} | {_cell(_cite(f.a))} | "
                     f"{_cell(_cite(f.b, f.also_b))} |" for f in similar)
    else:
        lines.append("None found.")
    lines.append("")

    for kind, label in (("only_in_a", a), ("only_in_b", b)):
        rows = [f for f in findings if f.type == kind]
        lines.extend([f"## Only in {label} ({len(rows)})", ""])
        lines.extend(f"- {_cite(f.a or f.b)}" for f in rows)
        lines.append("")

    aside = [r for r in requirements if r.excluded]
    lines.extend([f"## Drawing statements set aside ({len(aside)})", ""])
    if aside:
        lines.extend([
            "Read from a drawing by the local vision model and disputed by "
            "the second reading. They took no part in the comparison; read "
            "them from the sheet.", "",
            "| Location | First reading | Second reading |", "|---|---|---|"])
        lines.extend(f"| {_cell(r.ref)} (p.{r.page}) | {_cell(r.statement)} | "
                     f"{_cell(r.check_note)} |" for r in aside)
    else:
        lines.append("None.")
    lines.append("")
    return "\n".join(lines)


_HEAD_FILL = PatternFill("solid", fgColor="1F3864")
_SEV_FILL = {"high": "F8CBAD", "medium": "FFE699", "low": "E2EFDA"}


def _sheet(wb: Workbook, title: str, header: list[str],
           rows: list[list], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def excel_report(path: Path, meta_a: DocumentMeta, meta_b: DocumentMeta,
                 governing: DocumentMeta, findings: list[Finding],
                 reqs_a: list[Requirement], reqs_b: list[Requirement],
                 summary: dict) -> None:
    a, b = meta_a.label, meta_b.label
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in (
            ["Code comparison", f"{a} vs {b}"],
            [a, f"{meta_a.name} ({meta_a.filename})"],
            [b, f"{meta_b.name} ({meta_b.filename})"],
            ["Governing document", governing.label], [],
            [f"Requirements in {a}", summary["requirements_a"]],
            [f"Requirements in {b}", summary["requirements_b"]],
            ["Conflicts: direct contradictions", summary["contradictions"]],
            ["Conflicts: one code stricter", summary["more_stringent"]],
            ["Similar requirements", summary["similar"]],
            [f"Only in {a}", summary["unique_a"]],
            [f"Only in {b}", summary["unique_b"]],
            ["Flagged for engineer review", summary["needs_review"]],
            ["Drawing statements set aside",
             summary["set_aside_a"] + summary["set_aside_b"]]):
        ws.append(row)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70
    for c in ws["A"]:
        c.font = Font(bold=True)

    def side(r: Requirement | None) -> list:
        if r is None:
            return ["", "", "", ""]
        return [r.ref, r.page, "drawing" if r.from_figure else r.modality,
                r.statement]

    conflicts = [f for f in findings if f.category == "conflict"]
    ws = _sheet(
        wb, "Conflicts",
        ["ID", "Type", "Topic", "Severity",
         f"{a} clause", f"{a} page", f"{a} basis", f"{a} requirement",
         f"{b} clause", f"{b} page", f"{b} basis", f"{b} requirement",
         f"{b} also stated at", "Governing", "Stricter", "Action", "Why",
         "Drawing check", "Needs review", "Review note"],
        [[f.id, f.type.replace("_", " "), f.topic, f.severity, *side(f.a),
          *side(f.b), "; ".join(f.also_b), f.governing, f.stricter, f.action,
          f.rationale, f.drawing_check, "yes" if f.needs_review else "",
          f.review_reason] for f in conflicts],
        [8, 16, 24, 10, 18, 7, 10, 50, 22, 7, 10, 50, 30, 11, 10, 44, 50, 14,
         12, 44])
    for row, f in zip(ws.iter_rows(min_row=2), conflicts):
        if f.severity in _SEV_FILL:
            row[3].fill = PatternFill("solid", fgColor=_SEV_FILL[f.severity])

    similar = [f for f in findings if f.category == "similar"]
    _sheet(wb, "Similar",
           ["ID", f"{a} clause", f"{a} page", f"{a} basis",
            f"{a} requirement", f"{b} clause", f"{b} page", f"{b} basis",
            f"{b} requirement", f"{b} also stated at", "Similarity", "Why"],
           [[f.id, *side(f.a), *side(f.b), "; ".join(f.also_b), f.similarity,
             f.rationale] for f in similar],
           [8, 18, 7, 10, 55, 22, 7, 10, 55, 30, 10, 50])

    for kind, label in (("only_in_a", a), ("only_in_b", b)):
        rows = [f for f in findings if f.type == kind]
        _sheet(wb, f"Only in {label}"[:31],
               ["ID", "Clause", "Page", "Basis", "Requirement", "Subject",
                "Parameter", "Value"],
               [[f.id, *side(f.a or f.b), (f.a or f.b).subject,
                 (f.a or f.b).parameter, (f.a or f.b).value] for f in rows],
               [8, 22, 7, 10, 70, 26, 26, 30])

    for label, reqs in ((a, reqs_a), (b, reqs_b)):
        _sheet(wb, f"Requirements {label}"[:31],
               ["ID", "Clause", "Page", "Basis", "Requirement", "Subject",
                "Parameter", "Value", "Source text", "Quote verified",
                "References", "In comparison", "Drawing check", "Check note"],
               [[r.id, r.ref, r.page,
                 "drawing" if r.from_figure else r.modality, r.statement,
                 r.subject, r.parameter, r.value, r.quote,
                 "" if r.from_figure else ("yes" if r.quote_verified else "no"),
                 ", ".join(r.references), "no" if r.excluded else "yes",
                 r.drawing_check, r.check_note] for r in reqs],
               [14, 22, 7, 10, 70, 24, 24, 28, 60, 12, 26, 12, 14, 60])
    wb.save(path)
